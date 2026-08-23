"""Rich document parsing for Power Hub — PDF and Word text, tables, figures, images."""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)

MAX_CONTENT_CHARS = 400_000
MAX_PAGES = 80
MAX_IMAGES = 40


@dataclass
class ExtractedImage:
    name: str
    data: bytes
    page: int | None = None
    ocr_text: str = ""
    width: int | None = None
    height: int | None = None


@dataclass
class ParseResult:
    text: str = ""
    tables: list[str] = field(default_factory=list)
    figures: list[str] = field(default_factory=list)
    images: list[ExtractedImage] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def as_search_text(self) -> str:
        parts: list[str] = []
        if self.text.strip():
            parts.append(self.text.strip())
        for i, table in enumerate(self.tables, start=1):
            if table.strip():
                parts.append(f"[Table {i}]\n{table.strip()}")
        for i, figure in enumerate(self.figures, start=1):
            if figure.strip():
                parts.append(f"[Figure {i}]\n{figure.strip()}")
        for i, img in enumerate(self.images, start=1):
            caption = img.ocr_text.strip() or "(image with no readable text)"
            loc = f" page {img.page}" if img.page is not None else ""
            parts.append(f"[Image {i}{loc}: {img.name}]\n{caption}")
        joined = "\n\n".join(parts)
        return joined[:MAX_CONTENT_CHARS]


def _ocr_image_bytes(data: bytes) -> str:
    try:
        from PIL import Image
        import pytesseract

        image = Image.open(io.BytesIO(data))
        if image.mode not in ("RGB", "L"):
            image = image.convert("RGB")
        # Skip tiny decorative assets
        w, h = image.size
        if w < 40 or h < 40:
            return ""
        text = pytesseract.image_to_string(image) or ""
        return re.sub(r"[ \t]+\n", "\n", text).strip()
    except Exception as exc:
        logger.debug("OCR failed: %s", exc)
        return ""


def _table_to_text(rows: list[list[str | None]]) -> str:
    lines = []
    for row in rows:
        cells = [("" if c is None else str(c).strip().replace("\n", " ")) for c in row]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def parse_pdf(data: bytes) -> ParseResult:
    result = ParseResult()

    # --- Text + images via PyMuPDF ---
    try:
        import pymupdf

        doc = pymupdf.open(stream=data, filetype="pdf")
        text_parts: list[str] = []
        image_count = 0
        for page_index, page in enumerate(doc):
            if page_index >= MAX_PAGES:
                result.warnings.append(f"Truncated after {MAX_PAGES} pages")
                break
            page_text = page.get_text("text") or ""
            if page_text.strip():
                text_parts.append(f"--- Page {page_index + 1} ---\n{page_text.strip()}")

            # Captions / figure labels often appear near images as text blocks
            blocks = page.get_text("dict").get("blocks", [])
            for block in blocks:
                if block.get("type") != 0:
                    continue
                block_text = ""
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        block_text += span.get("text", "")
                    block_text += "\n"
                cleaned = block_text.strip()
                if re.match(r"^(figure|fig\.|table|chart|diagram)\b", cleaned, re.I):
                    result.figures.append(cleaned)

            for img_info in page.get_images(full=True):
                if image_count >= MAX_IMAGES:
                    break
                xref = img_info[0]
                try:
                    extracted = doc.extract_image(xref)
                except Exception:
                    continue
                img_bytes = extracted.get("image") or b""
                if not img_bytes:
                    continue
                ext = extracted.get("ext") or "png"
                name = f"page{page_index + 1}_img{image_count + 1}.{ext}"
                ocr = _ocr_image_bytes(img_bytes)
                result.images.append(
                    ExtractedImage(
                        name=name,
                        data=img_bytes,
                        page=page_index + 1,
                        ocr_text=ocr,
                        width=extracted.get("width"),
                        height=extracted.get("height"),
                    )
                )
                image_count += 1
        result.text = "\n\n".join(text_parts)
        doc.close()
    except Exception as exc:
        result.warnings.append(f"PDF text/image parse failed: {exc}")
        # Fallback: PyPDF2 text only
        try:
            from PyPDF2 import PdfReader

            reader = PdfReader(io.BytesIO(data))
            parts = []
            for i, page in enumerate(reader.pages[:MAX_PAGES]):
                parts.append(page.extract_text() or "")
            result.text = "\n".join(parts)
        except Exception as exc2:
            result.warnings.append(f"PDF fallback failed: {exc2}")

    # --- Tables via pdfplumber ---
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page_index, page in enumerate(pdf.pages[:MAX_PAGES]):
                try:
                    tables = page.extract_tables() or []
                except Exception:
                    tables = []
                for table in tables:
                    rendered = _table_to_text(table or [])
                    if rendered.strip():
                        result.tables.append(f"(page {page_index + 1})\n{rendered}")
    except Exception as exc:
        result.warnings.append(f"PDF table parse failed: {exc}")

    return result


def parse_docx(data: bytes) -> ParseResult:
    result = ParseResult()
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(io.BytesIO(data))

        # Body paragraphs + tables in document order
        body_parts: list[str] = []
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                para = Paragraph(child, document)
                text = para.text.strip()
                if not text:
                    continue
                style = (para.style.name if para.style is not None else "") or ""
                if re.match(r"^(figure|fig\.|table|caption)", text, re.I) or "Caption" in style:
                    result.figures.append(text)
                body_parts.append(text)
            elif child.tag == qn("w:tbl"):
                table = Table(child, document)
                rows = []
                for row in table.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
                rendered = _table_to_text(rows)
                if rendered.strip():
                    result.tables.append(rendered)

        result.text = "\n".join(body_parts)

        # Embedded images
        image_count = 0
        for rel in document.part.rels.values():
            if image_count >= MAX_IMAGES:
                break
            rel_type = getattr(rel, "reltype", "") or ""
            if "image" not in rel_type:
                continue
            try:
                img_bytes = rel.target_part.blob
            except Exception:
                continue
            if not img_bytes:
                continue
            content_type = getattr(rel.target_part, "content_type", "") or "image/png"
            ext = content_type.split("/")[-1].replace("jpeg", "jpg")
            if ext not in {"png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp", "emf", "wmf"}:
                ext = "bin"
            name = f"docx_img{image_count + 1}.{ext}"
            ocr = ""
            if ext in {"png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp"}:
                ocr = _ocr_image_bytes(img_bytes)
            result.images.append(
                ExtractedImage(name=name, data=img_bytes, ocr_text=ocr)
            )
            image_count += 1
    except Exception as exc:
        result.warnings.append(f"DOCX parse failed: {exc}")

    return result


def parse_doc_legacy(data: bytes, filename: str) -> ParseResult:
    """Best-effort for older .doc — extract readable strings."""
    result = ParseResult()
    result.warnings.append(f"Legacy .doc format has limited support for {filename}")
    try:
        # Prefer antiword/catdoc if present later; fall back to UTF-8/latin1 scrape
        text = data.decode("utf-8", errors="ignore")
        if len(text.strip()) < 40:
            text = data.decode("latin-1", errors="ignore")
        # Keep printable runs
        cleaned = re.sub(r"[^\x09\x0a\x0d\x20-\x7e\u00a0-\u024f]+", " ", text)
        cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
        result.text = cleaned.strip()[:MAX_CONTENT_CHARS]
    except Exception as exc:
        result.warnings.append(str(exc))
    return result


def parse_document(filename: str, data: bytes) -> ParseResult:
    ext = Path(filename).suffix.lstrip(".").lower()
    if ext == "pdf":
        return parse_pdf(data)
    if ext == "docx":
        return parse_docx(data)
    if ext == "doc":
        return parse_doc_legacy(data, filename)
    # Plain / other handled by caller
    result = ParseResult()
    try:
        result.text = data.decode("utf-8", errors="ignore")[:MAX_CONTENT_CHARS]
    except Exception:
        result.text = ""
    return result


def save_extracted_images(
    org_id: str,
    file_id: str,
    images: list[ExtractedImage],
    base_dir: Path | None = None,
) -> list[str]:
    """Persist extracted images next to vault files; return saved paths."""
    import config

    root = base_dir or Path(
        getattr(config, "POWERHUB_STORAGE_PATH", "./data/powerhub/files")
    ).parent / "extracted"
    dest = root / org_id / file_id
    dest.mkdir(parents=True, exist_ok=True)
    saved: list[str] = []
    for image in images:
        path = dest / image.name
        path.write_bytes(image.data)
        saved.append(str(path))
    return saved
